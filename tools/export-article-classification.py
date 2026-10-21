"""导出文章分类标注表（方案 B 简化版）。

用途：800 篇公众号文章只有标题可用（正文为空、微信原文有验证码拦），
让运营同事在 Excel 里把「不确定」的那部分标出来，再据此调准自动分类规则。

输出：docs/article-classification.xlsx
  Sheet1 标注       文章 ID / 标题 / 当前栏目 / 建议栏目(下拉) / 备注
  Sheet2 栏目对照    栏目 slug / 名称 / 当前篇数（下拉数据源）
  Sheet3 说明       怎么用

用法：python tools/export-article-classification.py
"""
import subprocess
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

MYSQL = r'C:\Environment\mysql-8.4.9\bin\mysql.exe'
OUT = Path(__file__).resolve().parent.parent / 'docs' / 'article-classification.xlsx'


def query(sql: str) -> list[list[str]]:
    result = subprocess.run(
        [MYSQL, '--login-path=daxin', '--host=127.0.0.1', '--port=3306', '--user=root',
         '--connect-timeout=5', '--batch', '--skip-column-names', '--default-character-set=utf8mb4',
         '--execute', f'USE daxin; {sql}'],
        capture_output=True, check=True,
    )
    text = result.stdout.decode('utf-8', errors='replace')
    return [line.split('\t') for line in text.splitlines() if line]


# 标题里可能有换行，SQL 里先替换掉，避免破坏按行解析
categories = query('SELECT slug, name, sort_order FROM article_categories ORDER BY sort_order, id;')
articles = query(
    "SELECT a.id, REPLACE(REPLACE(a.title, CHAR(10), ' '), CHAR(13), ' '), c.name, c.slug "
    'FROM articles a JOIN article_categories c ON c.id = a.category_id '
    "WHERE a.status = 'published' ORDER BY c.sort_order, a.id;"
)

counts = Counter(row[3] for row in articles)

wb = Workbook()

# ---------- Sheet1 标注 ----------
ws = wb.active
ws.title = '标注'
headers = ['文章 ID', '标题', '当前栏目', '建议栏目', '备注']
ws.append(headers)
header_fill = PatternFill('solid', fgColor='8F5F18')
for cell in ws[1]:
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')

for article_id, title, name, _slug in articles:
    ws.append([int(article_id), title, name, '', ''])

ws.freeze_panes = 'A2'
ws.column_dimensions['A'].width = 10
ws.column_dimensions['B'].width = 78
ws.column_dimensions['C'].width = 18
ws.column_dimensions['D'].width = 18
ws.column_dimensions['E'].width = 28
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    row[1].alignment = Alignment(wrap_text=True, vertical='top')
    for cell in row:
        cell.alignment = Alignment(vertical='top', wrap_text=cell is row[1])

# ---------- Sheet2 栏目对照 ----------
ws2 = wb.create_sheet('栏目对照')
ws2.append(['栏目 slug', '栏目名称', '当前篇数'])
for cell in ws2[1]:
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = header_fill
for slug, name, _order in categories:
    ws2.append([slug, name, counts.get(slug, 0)])
ws2.column_dimensions['A'].width = 24
ws2.column_dimensions['B'].width = 18
ws2.column_dimensions['C'].width = 12

# 建议栏目的下拉：引用栏目名称
last_row = max(ws.max_row, 2)
validation = DataValidation(
    type='list',
    formula1=f"'栏目对照'!$B$2:$B${len(categories) + 1}",
    allow_blank=True,
    showDropDown=False,
)
validation.error = '请从下拉里选择栏目'
validation.prompt = '留空表示「保持当前栏目」'
ws.add_data_validation(validation)
validation.add(f'D2:D{last_row}')

# ---------- Sheet3 说明 ----------
ws3 = wb.create_sheet('说明')
lines = [
    ('怎么用这张表', True),
    ('1. 只看「建议栏目」列。觉得当前栏目没问题的，留空即可（留空 = 不改）。', False),
    ('2. 觉得放错栏目的，从下拉里选正确栏目，必要时在「备注」写原因。', False),
    ('3. 重点看这几类：CPSC召回、劳动纠纷、货代案例、科普干货、普法干货。', False),
    ('4. 不认识 / 标题看不出内容的，留空并在备注写「不确定」，我会单独处理。', False),
    ('', False),
    ('为什么需要人工标注', True),
    ('· 800 篇文章只有标题，正文为空（采集器只存元数据）。', False),
    ('· 微信原文有验证码拦截，服务端抓不到正文（实测返回「环境异常」）。', False),
    ('· TRO侵权预警 / TRO成功案例 已按标题规则自动分好，本次主要校准其余栏目。', False),
    ('', False),
    ('当前各栏目篇数', True),
]
for slug, name, _order in categories:
    lines.append((f'  {name}（{slug}）：{counts.get(slug, 0)} 篇', False))
lines.append(('', False))
lines.append(('标完之后', True))
lines.append(('把文件放回 docs 目录，告诉我一声，我会按标注更新数据库，并把这些标注提炼成关键词规则。', False))

for text, bold in lines:
    ws3.append([text])
    if bold:
        ws3.cell(row=ws3.max_row, column=1).font = Font(bold=True, size=12)
ws3.column_dimensions['A'].width = 96

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f'已生成 {OUT}')
print(f'文章 {len(articles)} 篇，栏目 {len(categories)} 个')
for slug, name, _order in categories:
    print(f'  {name:<16} {counts.get(slug, 0)}')
